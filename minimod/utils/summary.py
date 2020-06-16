from tabulate import tabulate
import pandas as pd
import docx
import progressbar
import numpy as np

class PreOptimizationDataSummary:
    
    def __init__(self, 
                data, 
                benefit_col, 
                cost_col, 
                intervention_col, 
                space_col, 
                time_col, 
                benefit_title,
                cost_title = "Costs",
                intervention_title = 'Intervention',
                intervention_subset = None,
                space_title = 'Regions',
                time_title = 'Time'):
        
        self.data = data
        self.benefit_col = benefit_col
        self.cost_col = cost_col
        self.intervention_col = intervention_col
        self.space_col = space_col
        self.time_col = time_col
        self.benefit_title = benefit_title
        self.cost_title = cost_title
        self.intervention_title = intervention_title
        self.space_title = space_title
        self.time_title = time_title
        self.intervention_subset = intervention_subset
    
    def _rename_variables(self, data):
        
        df = (
            data
            .reset_index()
            .rename({
                self.benefit_col : self.benefit_title,
                self.cost_col : self.cost_title
            }, axis = 'columns')
        )
        
        if self.time_col in data.columns:
            df = df.rename({self.time_col : self.time_title}, axis = 'columns')
        if self.intervention_col in data.columns:
            df = df.rename({self.intervention_col : self.intervention_title}, axis = 'columns')
        if self.space_col in data.columns:
            df = df.rename({self.space_col : self.space_title}, axis = 'columns')
            
        return df
    
    def _create_national_intervention(self, data):
        
        subset_data = data[[self.intervention_col, self.space_col, self.time_col, self.benefit_col, self.cost_col]]
        
        national_sums = subset_data.groupby([self.intervention_col, self.time_col]).sum()
        
        national_sums[self.space_col] = 'National'
        
        df = (
            subset_data
            .set_index([self.intervention_col, self.space_col, self.time_col])
            .append(
                national_sums
                .reset_index()
                .set_index([self.intervention_col, self.space_col, self.time_col])
            )
            )
        
        return df
    
    def _construct_variables(self, data, variables_of_interest):
        
        for var in variables_of_interest:
            
            if variables_of_interest[var] == 'cost_per_benefit':
                
                # kwargs[var] = data[cost_col]/data[benefit_col]
                variables_of_interest[var] = lambda df: df[self.cost_col]/df[self.benefit_col]
            
            df = data.assign( var = variables_of_interest[var]).round({'var' : 2}).rename({'var' : var}, axis = 'columns')
            
        return df
    
    def _word_table(self, data, save_path):
        """Makes a word table out of a dataframe.
        Borrowed from stackoverflow:
        https://stackoverflow.com/questions/40596518/writing-a-python-pandas-dataframe-to-word-document
        """
        
        doc = docx.Document()
        
        table = doc.add_table(data.shape[0] + 1, data.shape[1])
        
        for j in range(data.shape[-1]):
            table.cell(0,j).text = data.columns[j]
        
        # add the rest of the data frame
        for i in progressbar.progressbar(range(data.shape[0])):
            for j in range(data.shape[-1]):
                table.cell(i+1,j).text = str(data.values[i,j])
                
        doc.save(save_path)
        print(f"Docx table saved in {save_path}")
                
    
    def _style(self, data, style = None, save_path = None):
        
        if style == 'html':
            return data.to_html()
        elif style == 'markdown':
            return data.to_markdown()
        elif style == 'latex':
            return data.to_latex()
        elif style == 'csv':
            data.reset_index().to_csv(save_path)
            print(f"CSV saved in {save_path}")
        elif style == 'word':
            self._word_table(data = data, save_path = save_path)
            print(f"Saved to {save_path}")
            
    def _num_formatter(self, data):
        
        data[self.benefit_col] = data[self.benefit_col].map("{:,.0f}".format)
        data[self.cost_col] = data[self.cost_col].map("${:,.0f}".format)
        
        return data
        
        
    def summary_table(self, 
                      variables_of_interest, 
                      grouping = None,
                      intervention_specific = True, 
                      style = None, 
                      save_path = 'table.docx'):
        
        if grouping == 'over_space':
            grouping = [self.space_col]
        elif grouping == 'over_time':
            grouping = [self.time_col]
            
            
        if intervention_specific:
            grouping.append(self.intervention_col)
            
        index_order = ['North', 'South', 'Cities', 'National']
        
        df = (
            self.data
            .pipe(self._create_national_intervention)
            .groupby(grouping).sum()
            .reset_index()
            .loc[lambda df: df[self.intervention_col].isin(self.intervention_subset)]
            .pipe(self._construct_variables, variables_of_interest)
            .pipe(self._num_formatter)
            .pipe(self._rename_variables)
            .drop(columns = 'index')
            )
                    
        if style is not None:
            self._style(data = df, style=style, save_path=save_path)
        
        return df
        
class OptimizationSummary:

    def __init__(self, model, table_fmt = "psql", decimals = 3, **kwargs):

        
        self.model = model
        self.table_fmt = table_fmt
        self.decimals = decimals

    def _print_specific_style(self, style, data):
        if style == "html":
            return self.html_print(data)
        elif style == "markdown":
            return self.markdown_print(data)
        elif style == "latex":
            return self.latex_print(data)
        elif style == 'pandas':
            return self.pandas_show(data)
        else:
            raise ValueError("style not available.")

    def html_print(self, data):
        return print(data.to_html())

    def markdown_print(self, data):
        return print(data.to_markdown())

    def latex_print(self, data):
        return print(data.to_latex())

    def pandas_show(self, data):
        return print(data)

    def _group_summarizer(self,
                          data = None,
                          intervention_specific=None,
                          over_time=True,
                          across_space=False,
                          style=None):

        if style is None:
            print("[Note]: style not specified, printing as pandas dataframe")
            print_style = self.pandas_show
        if style is not None:
            print_style = self._print_specific_style

        grouper = []

        if intervention_specific is None:
            intervention_specific = slice(None)

        if across_space:
            grouper.append(self.model.space_col)

        if over_time:
            grouper.append(self.model.time_col)

        summary_data = (data
                        .loc[(intervention_specific, slice(None), slice(None)), :]
                        .groupby(grouper)
                        .sum())

        return print_style(style, summary_data)
    
    def print_ratio(self, name, num, denom):
        
        try:
            ratio = num/denom
        except TypeError:
            ratio = "NaN"
            
        print(tabulate([(name, ratio)], tablefmt=self.table_fmt))
        
        return ratio
    
    def print_grouper(self, name, show_group = True, **kwargs):
        
        print(tabulate([(name, "")], tablefmt=self.table_fmt))
        
        if show_group:
            self._group_summarizer(**kwargs)    

    def print_generic(self, *args):

        for table in args:
            
            print(tabulate(table, tablefmt=self.table_fmt))
            
    def print_df(self, data):
        
        print(tabulate(data, headers = data.columns, tablefmt=self.table_fmt))
            
         
            
    
            
        
        
        
        
        


        
        



